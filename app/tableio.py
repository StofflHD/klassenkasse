"""Tabellen lesen und schreiben: CSV (Semikolon, UTF-8 mit BOM) und Excel (.xlsx).

Spalten werden als (Überschrift, Art) beschrieben; Art ist "text", "date",
"datetime", "euro" (Cent als int) oder "int". So bekommt Excel echte Zahlen und
Datumswerte, die CSV-Datei deutsche Schreibweisen.
"""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_ROWS = 2000  # Plausibilitätsgrenze beim Import

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _date(value):
    if not value:
        return None
    if isinstance(value, (date, datetime)):
        return value
    try:
        return datetime.fromisoformat(str(value).replace(" ", "T")) if len(str(value)) > 10 \
            else date.fromisoformat(str(value))
    except ValueError:
        return None


def euro_plain(cents):
    sign = "-" if cents < 0 else ""
    e, c = divmod(abs(int(cents)), 100)
    return f"{sign}{e:,}".replace(",", ".") + f",{c:02d}"


def csv_text(value):
    """Schutz vor CSV-/Formel-Injection in Tabellenkalkulationen."""
    value = "" if value is None else str(value)
    if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _csv_cell(kind, value):
    if value is None or value == "":
        return ""
    if kind == "euro":
        return euro_plain(value)
    if kind == "int":
        return str(value)
    if kind == "date":
        d = _date(value)
        return d.strftime("%d.%m.%Y") if d else csv_text(value)
    if kind == "datetime":
        d = _date(value)
        return d.strftime("%d.%m.%Y %H:%M") if isinstance(d, datetime) else csv_text(value)
    return csv_text(value)


def to_csv(columns, rows):
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow([h for h, _k in columns])
    for row in rows:
        w.writerow([_csv_cell(k, v) for (_h, k), v in zip(columns, row)])
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(columns, rows, title="Tabelle"):
    wb = Workbook()
    ws = wb.active
    # Excel verbietet []:*?/\ in Blattnamen und erlaubt höchstens 31 Zeichen
    ws.title = (re.sub(r"[\[\]:*?/\\]", "-", title).strip("' ") or "Tabelle")[:31]
    ws.append([h for h, _k in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([None] * len(columns))
        r = ws.max_row
        for i, ((_h, kind), value) in enumerate(zip(columns, row), start=1):
            if value is None or value == "":
                continue
            cell = ws.cell(row=r, column=i)
            if kind == "euro":
                cell.value = int(value) / 100
                cell.number_format = '#,##0.00 "€";-#,##0.00 "€"'
            elif kind == "int":
                cell.value = int(value)
            elif kind in ("date", "datetime"):
                d = _date(value)
                cell.value = d if d else str(value)
                if d:
                    cell.number_format = "DD.MM.YYYY HH:MM" if isinstance(d, datetime) else "DD.MM.YYYY"
            else:
                # immer als Text speichern – nie als Formel, auch wenn er mit = beginnt
                cell.value = str(value)
                cell.data_type = "s"
    for i, (header, kind) in enumerate(columns, start=1):
        width = max([len(header)] + [len(str(c.value or "")) for c in ws[get_column_letter(i)][1:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10 if kind != "text" else 8), 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- Lesen
def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # PLZ 80331.0 -> "80331"
    text = str(value).strip()
    if text.startswith("'") and text[1:2] in ("=", "+", "-", "@"):
        text = text[1:]  # Schutz-Apostroph aus unserem CSV-Export entfernen
    return text


def read_table(filename, data):
    """Datei (.xlsx oder .csv) -> (Überschriften, Zeilen als Listen von Text, Zeilennummern)."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Die Datei ist keine gültige Excel-Datei (.xlsx).")
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_text(v) for v in row])
            if len(rows) > MAX_ROWS + 1:
                raise ValueError(f"Zu viele Zeilen (höchstens {MAX_ROWS}).")
        wb.close()
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")  # ältere Excel-Versionen speichern CSV so
        sample = text[:4096]
        delimiter = max((";", ",", "\t"), key=sample.count)
        rows = [[_cell_text(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
        if len(rows) > MAX_ROWS + 1:
            raise ValueError(f"Zu viele Zeilen (höchstens {MAX_ROWS}).")
    else:
        raise ValueError("Bitte eine .xlsx- oder .csv-Datei wählen.")
    # Leerzeilen überspringen, aber die Zeilennummer der Datei behalten
    numbered = [(i, r) for i, r in enumerate(rows, start=1) if any(c for c in r)]
    if not numbered:
        raise ValueError("Die Datei ist leer.")
    header, body = numbered[0][1], numbered[1:]
    width = len(header)
    return header, [(r + [""] * width)[:width] for _i, r in body], [i for i, _r in body]
